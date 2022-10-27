from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

# Bar chart
from openpyxl.chart import BarChart, Reference

# Reference: 차트 만들 범위 참고
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()    # 차트 종류 설정
# bar_chart.add_data(bar_value)     # 차트 데이터 추가

# 만든 bar 차트를 ws에 추가
# ws.add_chart(bar_chart, "E1")
# ws.add_chart()

# Line chart
from openpyxl.chart import LineChart
# 제목을 포함하여 범위 지정
bar_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(bar_value, titles_from_data=True)

line_chart.title = "Course Scores"     # 제목 지정
line_chart.style = 20                  # 미리 정의된 스타일 중 20번 사용
line_chart.y_axis.title = "score"      # y축 제목 지정
line_chart.x_axis.title = "번호"       # x축 제목 지정

ws.add_chart(line_chart, "E2")

wb.save("sample_chart.xlsx")